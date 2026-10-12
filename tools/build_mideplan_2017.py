"""Extrae del libro «IDS 2017.xlsx» (MIDEPLAN, Índice de Desarrollo Social 2017) los datos de contraste:

  data/external/mideplan-ids-2017-distritos.csv   IDS 2017 por distrito + 5 dimensiones + área de desarrollo relativo
  data/external/mideplan-ids-2017-cantones.csv    IDS 2017 por cantón

Uso (desde la raíz del repositorio):  python tools/build_mideplan_2017.py "ruta/IDS 2017.xlsx"
Requiere: openpyxl. Es la edición que usó la tesis (2020). No es comparable con el IDS 2023.
"""
import csv
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
PROVINCES = {1: "San José", 2: "Alajuela", 3: "Cartago", 4: "Heredia", 5: "Guanacaste", 6: "Puntarenas", 7: "Limón"}
# Distritos con IDS imputado (no medido): informe oficial «Índice de Desarrollo Social 2017», Tabla 1 (p. 26).
# Se promedió o copió el valor de los distritos de procedencia o vecinos; no deben usarse para validar.
IMPUTED = {"10707": "Quitirrisí", "11912": "La Amistad", "20214": "San Lorenzo", "51105": "Matambú", "60206": "Caldera"}
# Distrito presente en el informe (Cuadro de posiciones, p. 30+) pero ausente del libro de Excel: Río Cuarto, cantón creado en 2017
# (en el informe usa el código 21601 «en lugar de 20306»). Valores tomados del informe (2 decimales).
FROM_PDF = {"21601": {"name": "Río Cuarto", "canton": "Río Cuarto", "ids": 44.75, "pos": 434, "area": "Muy bajo"}}
AREAS = {"Áreas de mayor desarrollo relativo": "Mayor", "Nivel medio": "Medio", "Nivel bajo": "Bajo", "Nivel muy bajo": "Muy bajo"}


def num(v):
    try:
        return float(v) if v not in (None, "") else None
    except ValueError:
        return None


def main(path: str) -> None:
    wb = openpyxl.load_workbook(path, data_only=True)

    ids = {}
    for r in wb["IDS segun PCD"].iter_rows(min_row=1, values_only=True):
        if r[0] and str(r[0]).isdigit():
            ids[str(r[0])] = {"name": str(r[1]).strip(), "ids": num(r[2]), "pos": int(float(r[3]))}

    dims = {}
    for r in wb["IDS dimensiones"].iter_rows(min_row=1, values_only=True):
        if r[2] and str(r[2]).isdigit():
            dims[str(r[2])] = {"canton": str(r[1]).strip(), "economico": num(r[4]), "participacion": num(r[5]), "salud": num(r[6]),
                               "educacion": num(r[7]), "seguridad": num(r[8]), "ids": num(r[9])}

    area, areas = None, {}
    for r in wb["IDS AMNDR"].iter_rows(min_row=1, values_only=True):
        a = r[0]
        if a and not str(a).isdigit() and r[1] is None:
            area = AREAS.get(str(a).strip(), area)
        elif a and str(a).isdigit():
            areas[str(a)] = area

    assert set(ids) == set(dims), (len(ids), len(dims))
    for c, x in FROM_PDF.items():  # completa el libro con el distrito que falta (ver FROM_PDF)
        assert c not in ids
        ids[c] = {"name": x["name"], "ids": x["ids"], "pos": x["pos"]}
        dims[c] = {"canton": x["canton"], "economico": None, "participacion": None, "salud": None, "educacion": None, "seguridad": None, "ids": x["ids"]}
        areas[c] = x["area"]
    assert set(IMPUTED) <= set(ids)
    bad = [c for c in ids if abs(ids[c]["ids"] - dims[c]["ids"]) > 0.05]
    print("distritos:", len(ids), "| imputados:", len(IMPUTED), "| IDS distinto entre hojas:", len(bad), "| con área:", len(areas))

    out = ROOT / "data" / "external"
    out.mkdir(parents=True, exist_ok=True)
    with open(out / "mideplan-ids-2017-distritos.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["codigo", "distrito", "canton", "provincia", "ids", "posicion", "area", "imputado", "salud", "participacion", "seguridad", "educacion", "economico"])
        rnd = lambda v: "" if v is None else round(v, 4)
        for c in sorted(ids, key=int):
            d = dims[c]
            w.writerow([c, ids[c]["name"], d["canton"], PROVINCES[int(c) // 10000], rnd(ids[c]["ids"]), ids[c]["pos"], areas.get(c, ""),
                        1 if c in IMPUTED else 0, rnd(d["salud"]), rnd(d["participacion"]), rnd(d["seguridad"]), rnd(d["educacion"]), rnd(d["economico"])])

    cantons = []
    for r in wb["IDS cantonal"].iter_rows(min_row=1, values_only=True):
        if r[0] and str(r[0]).isdigit() and num(r[2]) is not None:
            cantons.append((str(r[0]), str(r[1]).strip(), num(r[2])))
    with open(out / "mideplan-ids-2017-cantones.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["codigo", "canton", "provincia", "ids"])
        for c, n, v in sorted(cantons, key=lambda x: int(x[0])):
            w.writerow([c, n, PROVINCES[int(c) // 100], round(v, 4)])
    print("cantones:", len(cantons), "listo")


if __name__ == "__main__":
    main(sys.argv[1])
