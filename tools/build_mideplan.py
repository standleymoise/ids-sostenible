"""Extrae del libro «Tablas - IDS 2023.xlsx» (MIDEPLAN, Índice de Desarrollo Social 2023) los datos que usa
la web para validar y para elegir la zona:

  data/cr-distritos.json                       provincias > cantones > distritos con códigos oficiales (Registro Nacional, 2022)
  data/external/mideplan-ids-2023-distritos.csv  IDS 2023 por distrito + 5 dimensiones + quintil
  data/external/mideplan-ids-2023-cantones.csv   IDS 2023 por cantón + 5 dimensiones

Uso (desde la raíz del repositorio):  python tools/build_mideplan.py "ruta/Tablas - IDS 2023.xlsx"
Requiere: openpyxl. Fuente: MIDEPLAN, https://www.mideplan.go.cr/indice-desarrollo-social
"""
import csv
import json
import sys
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
PROVINCES = {1: "San José", 2: "Alajuela", 3: "Cartago", 4: "Heredia", 5: "Guanacaste", 6: "Puntarenas", 7: "Limón"}
REGION_SHEETS = ["Tabla 15", "Tabla 16", "Tabla 17", "Tabla 18", "Tabla 19", "Tabla 20"]


def norm(s) -> str:
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return " ".join("".join(c if c.isalnum() else " " for c in s).split())


def num(v):
    return float(v) if isinstance(v, (int, float)) else None


def main(path: str) -> None:
    wb = openpyxl.load_workbook(path, data_only=True)

    # --- división territorial oficial (Anexo 1b) ---
    divisions, cantons = {}, {}
    for r in wb["Anexo 1b"].iter_rows(min_row=1, values_only=True):
        if isinstance(r[0], (int, float)) and isinstance(r[2], (int, float)):
            divisions[int(r[2])] = {"canton_code": int(r[0]), "canton": str(r[1]).strip(), "district": str(r[3]).strip(), "region": str(r[4]).strip()}
            cantons[int(r[0])] = str(r[1]).strip()
    divisions.pop(60110, None)  # Isla del Coco: está en la división territorial pero MIDEPLAN no le calcula IDS
    assert len(cantons) == 84, len(cantons)
    assert len(divisions) == 490, len(divisions)

    tree = {}
    for code, d in sorted(divisions.items()):
        p = code // 10000
        tree.setdefault(p, {}).setdefault(d["canton_code"], {"code": d["canton_code"], "name": d["canton"], "districts": []})["districts"].append(
            {"code": code, "name": d["district"]})
    out = {
        "source": "Registro Nacional / IGN, División Territorial Administrativa 2022, tal como aparece en MIDEPLAN (Índice de Desarrollo Social 2023, Anexo 1).",
        "provinces": [{"code": p, "name": PROVINCES[p], "cantons": list(c.values())} for p, c in sorted(tree.items())],
    }
    (ROOT / "data" / "cr-distritos.json").write_text(json.dumps(out, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")

    # --- IDS distrital (Tabla 4) ---
    ids = {}
    for r in wb["Tabla 4"].iter_rows(min_row=1, values_only=True):
        if isinstance(r[0], (int, float)) and isinstance(r[2], (int, float)):
            ids[int(r[0])] = {"name": str(r[1]).strip(), "ids": float(r[2]), "pos": int(r[3])}
    assert len(ids) == 490 and set(ids) == set(divisions), (len(ids), set(divisions) ^ set(ids))

    # --- dimensiones por distrito (Tablas 15-20: cantón como encabezado, distritos debajo) ---
    key_to_code = {(norm(d["canton"]), norm(d["district"])): c for c, d in divisions.items()}
    dims, unmatched = {}, []
    for sheet in REGION_SHEETS:
        canton = None
        for r in wb[sheet].iter_rows(min_row=7, values_only=True):
            name = r[0]
            if not isinstance(name, str) or not name.strip():
                continue
            vals = [num(v) for v in r[1:7]]
            if all(v is None for v in vals):
                canton = name.strip()
                continue
            code = key_to_code.get((norm(canton), norm(name)))
            if code is None:
                unmatched.append((sheet, canton, name))
                continue
            dims[code] = {"salud": vals[0], "participacion": vals[1], "seguridad": vals[2], "educacion": vals[3], "economico": vals[4],
                          "ids_tabla": vals[5], "quintil": str(r[7]).strip()}
    if unmatched:
        print("SIN CRUZAR:", unmatched[:10], len(unmatched))
    missing = set(divisions) - set(dims)
    print("distritos con dimensiones:", len(dims), "faltan:", len(missing))
    bad = [c for c in dims if dims[c]["ids_tabla"] is not None and abs(dims[c]["ids_tabla"] - ids[c]["ids"]) > 1e-6]
    print("IDS de las tablas regionales difiere de la Tabla 4 en", len(bad), "distritos")

    outdir = ROOT / "data" / "external"
    outdir.mkdir(parents=True, exist_ok=True)
    with open(outdir / "mideplan-ids-2023-distritos.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["codigo", "distrito", "canton", "provincia", "ids", "posicion", "quintil", "salud", "participacion", "seguridad", "educacion", "economico"])
        for code in sorted(divisions):
            d, i, x = divisions[code], ids[code], dims.get(code, {})
            rnd = lambda v: "" if v is None else round(v, 4)
            w.writerow([code, d["district"], d["canton"], PROVINCES[code // 10000], rnd(i["ids"]), i["pos"], x.get("quintil", ""),
                        rnd(x.get("salud")), rnd(x.get("participacion")), rnd(x.get("seguridad")), rnd(x.get("educacion")), rnd(x.get("economico"))])

    # --- IDS cantonal (Tabla 21) ---
    rows = []
    for r in wb["Tabla 21"].iter_rows(min_row=1, values_only=True):
        if isinstance(r[0], (int, float)) and isinstance(r[1], str) and isinstance(r[8], (int, float)):
            rows.append(r)
    codes = {int(r[0]) for r in rows}
    print("cantones en Tabla 21:", len(rows), "faltan:", set(cantons) - codes)
    with open(outdir / "mideplan-ids-2023-cantones.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["codigo", "canton", "provincia", "ids", "salud", "participacion", "seguridad", "educacion", "economico"])
        for r in sorted(rows, key=lambda r: r[0]):
            w.writerow([int(r[0]), cantons.get(int(r[0]), r[1]), PROVINCES[int(r[0]) // 100], r[8], r[2], r[3], r[4], r[5], r[6]])
    print("listo")


if __name__ == "__main__":
    main(sys.argv[1])
